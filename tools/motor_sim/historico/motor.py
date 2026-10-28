"""Replica fiel do motor atual (StructureEngine + ExercisePreFilter + SlotFiller)."""
import openpyxl, random, collections, itertools, statistics

# ---------- catalogo: as 79 da lista principal, pos-V64 + regras ratificadas ----------
TARGETS = ["CHEST","BACK","SHOULDERS","BICEPS","TRICEPS","FOREARMS","QUADS","POSTERIOR","CALVES","GLUTES","CORE"]
MUSCLEGROUP = {"CHEST","BACK","BICEPS","TRICEPS","FOREARMS","SHOULDERS","LEGS","GLUTES","CORE"}  # sem CALVES/TRAPEZIUS/LOWER_BACK

def load():
    wb = openpyxl.load_workbook('/mnt/user-data/uploads/FitJourney/tools/cache/lista-principal-detalhada.xlsx', data_only=True)
    ws = wb['Lista']; hdr=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
    out=[]
    for r in ws.iter_rows(min_row=2, values_only=True):
        d=dict(zip(hdr,r))
        nm=d.get('Exercício (PT)')
        if not nm: continue
        prim=[p.strip() for p in str(d.get('Músculos primários') or '').split(',') if p.strip()]
        ex = dict(
            n=int(d['Nº']), name=nm, cat=str(d.get('Categoria') or ''),
            eq=str(d.get('Equipamento') or ''),
            pat=str(d.get('Padrão de movimento (no motor)') or 'NONE'),
            comp=(str(d.get('Composto')).strip().lower()=='sim'),
            prim=prim, level=str(d.get('Nível') or ''),
            base=(str(d.get('is_base (atual)')).strip().lower()=='sim'),
        )
        out.append(ex)
    # regras ratificadas 2026-10-10 / V64
    for ex in out:
        if 'Rosca Inversa' in ex['name']:
            ex['cat']='FOREARMS'; ex['prim']=['FOREARMS']
        if 'Adutora' in ex['name']:
            ex['prim']=['GLUTES']
        if 'Face Pull' in ex['name']:
            ex['cat']='BACK'; ex['prim']=['BACK']
        if 'Remada Alta' in ex['name']:
            ex['cat']='BACK'; ex['pat']='VERTICAL_PULL'
            if 'Halteres' in ex['name']: ex['eq']='DUMBBELL'   # correcao de equipamento prevista na V64
        # toExercise(): enum invalido cai em silencio
        ex['prim']=[p for p in ex['prim'] if p in MUSCLEGROUP]
    return out

CAT = load()

# ---------- VolumeModel ----------
V = {  # ini, inter, adv, mrv
 "CHEST":(8,14,18,22),"BACK":(10,16,20,25),"SHOULDERS":(6,12,16,22),
 "BICEPS":(5,9,12,16),"TRICEPS":(6,10,14,18),"FOREARMS":(3,6,9,12),
 "QUADS":(6,12,16,20),"POSTERIOR":(6,10,14,18),"CALVES":(5,9,12,16),
 "GLUTES":(4,8,12,16),"CORE":(5,8,10,16)}
def weekly_sets(m, level, goal):
    v=V[m]
    if goal=="GENERAL_HEALTH": return v[0]
    return {"BEGINNER":v[0],"INTERMEDIATE":v[1],"ADVANCED":v[2]}[level]
def mrv(m): return V[m][3]
def age_band(a):
    if a is None: return "YOUNG"
    return "SENIOR" if a>=60 else ("ADULT" if a>=40 else "YOUNG")
def age_factor(a): return {"YOUNG":1.0,"ADULT":0.9,"SENIOR":0.75}[age_band(a)]
FOCUS_TARGETS = {"CHEST":["CHEST"],"BACK":["BACK"],"SHOULDERS":["SHOULDERS"],"BICEPS":["BICEPS"],
 "TRICEPS":["TRICEPS"],"FOREARMS":["FOREARMS"],"LEGS":["QUADS","POSTERIOR","CALVES"],
 "GLUTES":["GLUTES"],"CORE":["CORE"]}
ROLE_P = {  # (reps, rest, rir) normal / conservador
 "COMPOSTO_PESADO":(("5-8",150,2),("6-10",150,3)),
 "COMPOSTO_ACESSORIO":(("8-12",105,1),("8-12",105,3)),
 "ISOLAMENTO":(("12-15",75,1),("10-15",75,3))}
def params_for(role, level, goal, age):
    cons = goal=="GENERAL_HEALTH" or level=="BEGINNER" or age_band(age)=="SENIOR"
    return ROLE_P[role][1 if cons else 0]

# ---------- splits ----------
UPPER={"CHEST","BACK","SHOULDERS","BICEPS","TRICEPS","FOREARMS"}
LOWER={"QUADS","POSTERIOR","CALVES","GLUTES","CORE"}
PUSH={"CHEST","SHOULDERS","TRICEPS"}; PULL={"BACK","BICEPS","FOREARMS"}
LEGS={"QUADS","POSTERIOR","CALVES","GLUTES"}
CHEST_BACK={"CHEST","BACK"}; SHOULDERS_ARMS={"SHOULDERS","BICEPS","TRICEPS","FOREARMS"}
FULL=set(TARGETS)
SPLIT_OPTS={2:["FULL_BODY"],3:["FULL_BODY","UPPER_LOWER_FULL","PUSH_PULL_LEGS"],
 4:["UPPER_LOWER","FULL_BODY"],5:["UL_PPL"],6:["PUSH_PULL_LEGS","ARNOLD"]}
def resolve_split(days, chosen):
    opts=SPLIT_OPTS[max(2,min(6,days))]
    return chosen if chosen in opts else opts[0]
def cycle(days, base):
    n=len(base); rounds=(days+n-1)//n; out=[]
    for i in range(days):
        lab,mus=base[i%n]
        suf=" "+chr(ord('A')+i//n) if rounds>1 else ""
        out.append((lab+suf,mus))
    return out
def split_template(days, split):
    if split=="FULL_BODY": return cycle(days,[("Full Body",FULL)])
    if split=="UPPER_LOWER": return cycle(days,[("Upper",UPPER),("Lower",LOWER)])
    if split=="UPPER_LOWER_FULL": return [("Upper",UPPER),("Lower",LOWER),("Full Body",FULL)]
    if split=="PUSH_PULL_LEGS": return cycle(days,[("Push",PUSH),("Pull",PULL),("Legs",LEGS)])
    if split=="UL_PPL": return [("Upper",UPPER),("Lower",LOWER),("Push",PUSH),("Pull",PULL),("Legs",LEGS)]
    if split=="ARNOLD": return cycle(days,[("Peito+Costas",CHEST_BACK),("Pernas",LEGS),("Ombros+Braços",SHOULDERS_ARMS)])

CANON=["CHEST","BACK","SHOULDERS","QUADS","POSTERIOR","GLUTES","BICEPS","TRICEPS","FOREARMS","CALVES","CORE"]
BIGS=["CHEST","BACK","QUADS","POSTERIOR","SHOULDERS"]
SMALLS=["GLUTES","BICEPS","TRICEPS","FOREARMS","CALVES","CORE"]
MIN_SETS=3; MAX_SETS_EX=5; SESSION_MAX_SETS=8; MAX_EX=6; MAX_SETS_TOTAL=20; SETS_PER_EX=3; FOCUS_BONUS=3

def rnd(x):  # Kotlin roundToInt: meio pra cima
    import math
    return math.floor(x+0.5)

def rotate(lst, by):
    if not lst: return lst
    n=((by%len(lst))+len(lst))%len(lst)
    return lst[n:]+lst[:n]

def trained_muscles(muscles, day_index, focus):
    if len(muscles)<=MAX_EX: return [m for m in CANON if m in muscles]
    bigs=[m for m in BIGS if m in muscles]; smalls=[m for m in SMALLS if m in muscles]
    budget=max(0,MAX_EX-len(bigs))
    fs=[m for m in smalls if m in focus]; rs=[m for m in smalls if m not in focus]
    kept_f=rotate(fs,day_index)[:budget]
    rem=max(0,budget-len(kept_f)); kept_r=rotate(rs,day_index)[:rem]
    kept=set(bigs)|set(kept_f)|set(kept_r)
    return [m for m in CANON if m in kept]

def allocate_exercises(trained, session_sets):
    natural={m:max(1,rnd(session_sets[m]/SETS_PER_EX)) for m in trained}
    alloc={m:1 for m in trained}; used=len(trained)
    while used<MAX_EX:
        cands=[m for m in trained if alloc[m]<natural[m]]
        if not cands: break
        c=max(cands,key=lambda m:(session_sets[m]/alloc[m]))
        alloc[c]+=1; used+=1
    return alloc

ISO_ONLY={"CALVES","CORE","BICEPS","TRICEPS","FOREARMS"}
def role_for(target,i):
    if target in ISO_ONLY: return "ISOLAMENTO"
    return ["COMPOSTO_PESADO","COMPOSTO_ACESSORIO"][i] if i<2 else "ISOLAMENTO"
ROLE_ORD={"COMPOSTO_PESADO":0,"COMPOSTO_ACESSORIO":1,"ISOLAMENTO":2}

def build_skeleton(goal, level, days, focus_groups, split_pref=None, age=None):
    days=max(2,min(6,days)); split=resolve_split(days,split_pref)
    af=age_factor(age)
    ft=set(itertools.chain.from_iterable(FOCUS_TARGETS[g] for g in focus_groups))
    weekly={m:min(rnd(weekly_sets(m,level,goal)*af)+(FOCUS_BONUS if m in ft else 0), mrv(m)) for m in TARGETS}
    tmpl=split_template(days,split)
    freq={m:max(1,sum(1 for _,mus in tmpl if m in mus)) for m in TARGETS}
    day_sk=[]
    for di,(label,mus) in enumerate(tmpl):
        trained=trained_muscles(mus,di,ft)
        ss={m:max(MIN_SETS,min(SESSION_MAX_SETS,rnd(weekly[m]/freq[m]))) for m in trained}
        ec=allocate_exercises(trained,ss)
        sp={m:max(MIN_SETS,min(MAX_SETS_EX,rnd(ss[m]/ec[m]))) for m in trained}
        def total(): return sum(ec[m]*sp[m] for m in trained)
        while total()>MAX_SETS_TOTAL:
            c=[m for m in trained if sp[m]>MIN_SETS]
            if not c: break
            a=max(c,key=lambda m:(sp[m],-CANON.index(m)))
            sp[a]-=1
        slots=[]
        for m in trained:
            for i in range(ec[m]):
                role=role_for(m,i); reps,rest,rir=params_for(role,level,goal,age)
                slots.append(dict(target=m,role=role,sets=sp[m],reps=reps,rest=rest,rir=rir))
        slots.sort(key=lambda s:ROLE_ORD[s['role']])
        day_sk.append(dict(label=label,slots=slots))
    return dict(days=day_sk,split=split,weekly=weekly,freq=freq)

# ---------- pool ----------
HOME_EQ={"DUMBBELL","BODYWEIGHT","BAND"}
def levels_up_to(l): return {"BEGINNER":["BEGINNER"],"INTERMEDIATE":["BEGINNER","INTERMEDIATE"],
                             "ADVANCED":["BEGINNER","INTERMEDIATE","ADVANCED"]}[l]
def pool_for(env, level):
    al=levels_up_to(level)
    return [e for e in CAT if e['level'] in al and (env=="ACADEMIA" or e['eq'] in HOME_EQ)]

QUAD_PAT={"SQUAT","LUNGE","KNEE_EXTENSION"}; POST_PAT={"HINGE","KNEE_FLEXION"}
def matches(e,t):
    if t=="CHEST": return "CHEST" in e['prim']
    if t=="BACK": return "BACK" in e['prim']
    if t=="SHOULDERS": return "SHOULDERS" in e['prim']
    if t in ("BICEPS","TRICEPS","FOREARMS"): return e['cat']==t
    if t=="QUADS": return e['pat'] in QUAD_PAT
    if t=="POSTERIOR": return e['pat'] in POST_PAT and "LEGS" in e['prim']
    if t=="CALVES": return e['cat']=="CALVES"
    if t=="GLUTES": return "GLUTES" in e['prim']
    if t=="CORE": return "CORE" in e['prim'] or e['cat']=="CORE"
EQW={"BARBELL":1.0,"MACHINE":1.0,"CABLE":1.0,"DUMBBELL":0.8,"KETTLEBELL":0.8,
     "BAND":0.4,"GYMSTICK":0.4,"BODYWEIGHT":0.2,"IMPROVISED":0.2}
def fill_day(day, pool, focus_names, level, used, rng):
    out=[]; used_pat=set()
    for s in day['slots']:
        cand=[e for e in pool if matches(e,s['target'])]
        cand=[e for e in cand if (e['comp'] if s['role']!="ISOLAMENTO" else not e['comp'])]
        cand=[e for e in cand if e['n'] not in used]
        base=[e for e in cand if e['base']]
        cand=base if base else cand
        if not cand:
            out.append(dict(slot=s,ex=None)); continue
        def sc(e):
            v=0.0
            if any(p in focus_names for p in e['prim']): v+=4.0
            if e['level']==level: v+=2.0
            if s['role']!="ISOLAMENTO" and e['comp']: v+=2.0
            v+=EQW.get(e['eq'],0.5)
            if e['pat'] in used_pat: v-=2.5
            return v+rng.random()*0.35
        best=max(cand,key=sc)
        used.add(best['n']); used_pat.add(best['pat'])
        out.append(dict(slot=s,ex=best))
    return out

def generate(goal,level,days,focus_groups,env,age,split_pref=None,seed=0):
    sk=build_skeleton(goal,level,days,focus_groups,split_pref,age)
    pool=pool_for(env,level)
    rng=random.Random(seed); used=set(); prog=[]
    fn=set(focus_groups)
    for d in sk['days']:
        prog.append(dict(label=d['label'], filled=fill_day(d,pool,fn,level,used,rng)))
    return sk,prog
